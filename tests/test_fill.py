import asyncio
import io
import sqlite3
import zipfile
from pathlib import Path

import docx
import pytest
import yaml
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QLineEdit

from blankey.core import RequestKind
from blankey.mcp.server import build_server
from blankey.templates import TemplateKind, engine, fill
from blankey.templates import docx as docx_template
from blankey.ui.fill import FillDialog
from blankey.vault import FieldInput, FieldType

REPLACEMENTS = [  # applied in order: longer dot blanks first
    {"find": "...........2026", "var": "contract_date"},
    {"find": "„ACME HOLDINGS“ ЕООД", "var": "company_name"},
    {"find": "999000111", "var": "company_eik"},
    {"find": "John Smith", "var": "manager_name"},
    {"find": "..........", "var": "manager_egn", "occurrence": 1},
    {"find": "..........", "var": "manager_id_card", "occurrence": 1},
    {"find": "34 000", "var": "loan_amount"},
    {"find": ".........", "var": "contract_no"},
]
FIELDS = {
    "company_name": "{{ company.name }}",
    "company_eik": "{{ company.eik }}",
    "manager_name": "{{ manager.name }}",
    "manager_egn": "{{ manager.egn }}",
    "contract_date": "{{ today() }}",
}
ROLES = {"company": "lender", "manager": "borrower"}


def text_of(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


@pytest.fixture
def template(app, filled_docx, tmp_path):
    target = tmp_path / "template.docx"
    docx_template.tokenize(filled_docx, target, REPLACEMENTS)
    return app.templates.save(
        "loan", "Loan", TemplateKind.DOCX, ROLES, FIELDS, source_file=target, labels={"loan_amount": "Сума"}
    )


@pytest.fixture
def profiles(app):
    company = app.vault.create_profile("Delta", "company")
    app.vault.set_values(company, [FieldInput("name", "„ACME“ ЕООД"), FieldInput("eik", "999000111")])
    manager = app.vault.create_profile("John", "person")
    app.vault.set_values(
        manager, [FieldInput("name", "Тест Тестов"), FieldInput("egn", "7501020018", type=FieldType.EGN)]
    )
    return {"company": company, "manager": manager}


def test_tokenize_and_variables(template):
    text = text_of(template.source)
    assert "{{ company_name }}, ЕИК {{ company_eik }}" in text
    assert "ЕГН {{ manager_egn }}, л.к. № {{ manager_id_card }}" in text
    assert "Договор № {{ contract_no }} / {{ contract_date }} г." in text
    assert docx_template.variables(template.source_path) == [
        "contract_no",
        "contract_date",
        "company_name",
        "company_eik",
        "manager_name",
        "manager_egn",
        "manager_id_card",
        "loan_amount",
    ]
    contexts = docx_template.blank_contexts(template.source_path)
    assert contexts["loan_amount"] == ["Заем в размер на [____] евро."]
    assert len(contexts["manager_name"]) == 2  # paragraph and table cell


def test_tokenize_errors(filled_docx, tmp_path):
    with pytest.raises(ValueError, match="spans several runs"):
        docx_template.tokenize(filled_docx, tmp_path / "x.docx", [{"find": "ЕООД, ЕИК", "var": "x"}])
    with pytest.raises(ValueError, match="not found"):
        docx_template.tokenize(filled_docx, tmp_path / "x.docx", [{"find": "nope", "var": "x"}])
    with pytest.raises(ValueError, match="no occurrence 9"):
        docx_template.tokenize(filled_docx, tmp_path / "x.docx", [{"find": "John", "var": "x", "occurrence": 9}])


def test_describe_warns_about_unbound_tags(template):
    warnings = engine.describe_template(template)["warnings"]
    assert "Tag without binding: loan_amount" in warnings


def test_fill_set_encrypted_round_trip(app):
    bindings = {"manager_egn": {"path": "manager.egn"}, "secret": {"value": "Тайна Стойност"}}
    fill_set_id = app.vault.save_fill_set("s", ["loan"], {"manager": 1}, bindings)
    raw = sqlite3.connect(app.config.db_path).execute("SELECT ciphertext FROM fill_sets").fetchone()[0]
    assert "Тайна".encode() not in raw
    info, loaded = app.vault.get_fill_set(fill_set_id)
    assert (info.name, info.templates, info.profiles, loaded) == ("s", ["loan"], {"manager": 1}, bindings)
    app.vault.save_fill_set("s2", ["loan"], {}, {}, fill_set_id)
    assert [f.name for f in app.vault.list_fill_sets()] == ["s2"]


def test_generate_fill_mixes_paths_values_and_defaults(app, template, profiles):
    bindings = {
        "manager_name": {"path": "manager.name"},
        "loan_amount": {"value": "34 000"},
        "contract_no": {"value": ""},
        "manager_id_card": {"path": "manager.egn"},
    }
    fill_set_id = app.vault.save_fill_set("delta", ["loan"], profiles, bindings)
    [doc_id] = app.generate_fill(fill_set_id)
    text = text_of(app.vault.load_document(doc_id))
    assert "„ACME“ ЕООД, ЕИК 999000111, представлявано от Тест Тестов" in text
    assert "ЕГН 7501020018, л.к. № 7501020018" in text
    assert "34 000 евро" in text
    assert "None" not in text and "{{" not in text


def test_export_import_and_bundle(app, template, profiles, tmp_path):
    fill_set_id = app.vault.save_fill_set("delta", ["loan"], profiles, {"loan_amount": {"value": "1"}})
    exported = yaml.safe_load(app.export_fill_set(fill_set_id))
    assert exported["bindings"] == {"loan_amount": {"value": "1"}}
    exported["bindings"]["loan_amount"]["value"] = "2"
    assert app.import_fill_set(yaml.safe_dump(exported, allow_unicode=True)) == fill_set_id
    assert app.vault.get_fill_set(fill_set_id)[1]["loan_amount"] == {"value": "2"}
    app.export_fill_bundle(fill_set_id, tmp_path / "b.zip")
    assert sorted(zipfile.ZipFile(tmp_path / "b.zip").namelist()) == [
        "fill.yaml",
        "loan/manifest.yaml",
        "loan/template.docx",
    ]


def test_default_binding():
    assert fill.default_binding(fill.FillVar("x", defaults={"t": "{{ manager.egn }}"}), {"manager"}) == {
        "path": "manager.egn"
    }
    assert fill.default_binding(fill.FillVar("x", defaults={"t": "{{ manager.name | upper }}"}), {"manager"}) == {
        "default": True
    }
    assert fill.default_binding(fill.FillVar("x"), {"manager"}) == {"value": ""}


def test_request_fill_validates_and_hides_values(qtbot, app, template, profiles, filled_docx, opened_urls):
    mcp = build_server(app)

    def call(tool, **arguments):
        return asyncio.run(mcp.call_tool(tool, arguments))

    assert "ACME" in call("inspect_docx", path=str(filled_docx)).model_dump_json(ensure_ascii=False)
    with pytest.raises(Exception, match="Missing profiles"):
        call("request_fill", templates=["loan"], profiles={"company": profiles["company"]}, name="x")
    with pytest.raises(Exception, match="no vault field"):
        call("request_fill", templates=["loan"], profiles=profiles, name="x", paths={"manager_egn": "manager.nope"})
    with pytest.raises(Exception, match="Unknown variable"):
        call("request_fill", templates=["loan"], profiles=profiles, name="x", values={"nope": "1"})
    call(
        "request_fill",
        templates=["loan"],
        profiles=profiles,
        name="delta",
        values={"loan_amount": "34 000"},
        paths={"manager_id_card": "manager.egn"},
    )
    request = app.vault.list_requests("pending")[0]
    assert request.kind == RequestKind.FILL

    dialog = FillDialog(
        app,
        request.payload["templates"],
        request.payload["profiles"],
        request.payload["name"],
        request.payload["values"],
        request.payload["paths"],
        request=request,
    )
    qtbot.addWidget(dialog)
    bindings = dialog.current_bindings()
    assert bindings["loan_amount"] == {"value": "34 000"}
    assert bindings["manager_id_card"] == {"path": "manager.egn"}
    assert bindings["manager_egn"] == {"path": "manager.egn"}
    assert bindings["contract_date"] == {"default": True}
    row = list(dialog.vars).index("contract_no")
    dialog.set_binding(row, {"value": "17"})
    dialog._save(generate=True)

    result = app.vault.get_request(request.id).result
    assert result["outcome"] == "generated" and len(result["document_ids"]) == 1
    assert result["empty"] == []
    assert len(opened_urls) == 1 and opened_urls[0].endswith(".docx")
    assert "Договор № 17" in text_of(Path(opened_urls[0]).read_bytes())
    assert "Договор № 17" in text_of(app.vault.load_document(result["document_ids"][0]))
    listed = call("list_fill_sets").model_dump_json()
    assert "7501020018" not in listed and "Тест" not in listed and '"loan_amount":"value"' in listed.replace(" ", "")


def test_fill_dialog_cancel_and_save(qtbot, app, template, profiles):
    request_id = app.create_request(
        RequestKind.FILL, {"templates": ["loan"], "profiles": profiles, "name": "n", "values": {}, "paths": {}}
    )
    dialog = FillDialog(app, ["loan"], profiles, "n", request=app.vault.get_request(request_id))
    qtbot.addWidget(dialog)
    dialog._save(generate=False)
    dialog.reject()
    result = app.vault.get_request(request_id).result
    assert result["outcome"] == "saved"
    assert "loan_amount" in result["empty"]
    reopened = FillDialog(app, [], {}, "n")
    qtbot.addWidget(reopened)
    assert reopened.fill_set_id == result["fill_set_id"]


def test_source_menu_and_typing(qtbot, app, template, profiles):
    dialog = FillDialog(app, ["loan"], profiles, "menu", values={"loan_amount": "34 000"})
    qtbot.addWidget(dialog)
    dialog.show()
    row = list(dialog.vars).index("loan_amount")
    actions = [a for a in dialog.source_menu(row).actions() if not a.isSeparator()]
    texts = [a.text() for a in actions]
    assert texts[:2] == ["Type a value…", "Claude  —  34 000"]
    assert [a.text() for a in actions if a.isChecked()] == ["Claude  —  34 000"]
    assert "manager.egn  —  7501020018" in texts

    next(a for a in actions if a.text().startswith("manager.egn")).trigger()
    assert dialog.current_bindings()["loan_amount"] == {"path": "manager.egn"}

    dialog.start_typing(row)
    editor = dialog.table.findChild(QLineEdit)
    assert editor is not None and editor.text() == ""
    editor.setText("50 000")
    qtbot.keyClick(editor, Qt.Key.Key_Return)
    qtbot.waitUntil(lambda: dialog.current_bindings()["loan_amount"] == {"value": "50 000"})
    assert dialog.describe_row(row)[:3] == ("Value", "", "50 000")


def test_docx_to_html_shows_blanks_as_chips(template):
    html = docx_template.to_html(template.source_path, {"loan_amount": "Сума & <лихва>"})
    link = '<a href="field:loan_amount" name="blank-loan_amount">'
    assert link + '<span class="blank b-loan_amount">&nbsp;Сума &amp; &lt;лихва&gt;&nbsp;</span></a>' in html
    assert '<span class="blank b-company_name">&nbsp;company_name&nbsp;</span>' in html
    assert "<b>1. " in html
    assert "<table" in html and "<td>" in html
    assert "{{" not in html

    filled = docx_template.to_html(template.source_path, {}, {"loan_amount": "34 000 & 5%", "company_name": " "})
    assert '<span class="blank b-loan_amount">&nbsp;34 000 &amp; 5%&nbsp;</span>' in filled
    assert '<span class="blank b-company_name empty">&nbsp;company_name&nbsp;</span>' in filled


def _pdf_template(app, form_pdf, tmp_path):
    source = tmp_path / "form.pdf"
    source.write_bytes(form_pdf)
    app.templates.save("pdf", "PDF", TemplateKind.PDF_FORM, {"a": "x"}, {"name": "{{ a.name }}"}, source_file=source)


def _select(dialog, var: str) -> None:
    dialog.table.setCurrentCell(list(dialog.vars).index(var), 0)


def test_preview_highlights_the_selected_field_on_the_pdf(qtbot, app, form_pdf, tmp_path):
    _pdf_template(app, form_pdf, tmp_path)
    [var] = fill.variables([app.templates.get("pdf")]).values()
    assert var.label == "Name"
    dialog = FillDialog(app, ["pdf"], {}, "pdf-fill")
    qtbot.addWidget(dialog)
    dialog.show()
    page = dialog.preview.pages[0]
    [view] = [v for v in page.views if "name" in v.rects]
    assert view.active == "name"  # the first row is selected when the form opens
    center = view.field_rects("name")[0].center().toPoint()
    image = view.grab().toImage()
    ratio = image.devicePixelRatio()
    pixel = image.pixelColor(int(center.x() * ratio), int(center.y() * ratio))
    accent = dialog.palette().color(dialog.palette().ColorRole.Accent)
    assert pixel != Qt.GlobalColor.white and abs(pixel.hue() - accent.hue()) < 20


def test_preview_highlights_docx_blanks_and_switches_documents(qtbot, app, template, form_pdf, tmp_path):
    _pdf_template(app, form_pdf, tmp_path)
    dialog = FillDialog(app, ["loan", "pdf"], {}, "both")
    qtbot.addWidget(dialog)
    dialog.show()
    docx_page, pdf_page = dialog.preview.pages
    assert dialog.preview.tabs.isVisible() and dialog.preview.tabs.count() == 2
    _select(dialog, "loan_amount")
    assert dialog.preview.stack.currentWidget() is docx_page
    assert ".b-loan_amount" in docx_page.browser.document().defaultStyleSheet()
    _select(dialog, "name")
    assert dialog.preview.stack.currentWidget() is pdf_page and dialog.preview.tabs.currentIndex() == 1


def test_preview_finds_typst_values_on_the_page(qtbot, app):
    source = "#let data = json(bytes(sys.inputs.data))\n= Декларация\nДолуподписаният #data.name, ЕГН #data.egn\n"
    app.templates.save(
        "decl",
        "Декларация",
        TemplateKind.TYPST,
        {"applicant": "person"},
        {"name": "{{ applicant.name }}", "egn": "{{ applicant.egn }}"},
        source_text=source,
    )
    dialog = FillDialog(app, ["decl"], {}, "typst")
    qtbot.addWidget(dialog)
    dialog.show()
    [page] = dialog.preview.pages
    assert page.names == {"name", "egn"}
    _select(dialog, "egn")
    [view] = page.views
    assert view.active == "egn" and view.field_rects("egn")[0].width() > 0


def test_preview_shows_values_and_clicks_select_the_field(qtbot, app, template, profiles, form_pdf, tmp_path):
    from PySide6.QtCore import QPoint, QUrl

    _pdf_template(app, form_pdf, tmp_path)
    dialog = FillDialog(app, ["loan", "pdf"], profiles, "live", values={"loan_amount": "34 000"})
    qtbot.addWidget(dialog)
    dialog.show()
    docx_page, pdf_page = dialog.preview.pages
    assert "34 000" in docx_page.browser.toPlainText()
    dialog.set_binding(list(dialog.vars).index("loan_amount"), {"value": "50 000"})
    qtbot.waitUntil(lambda: "50 000" in docx_page.browser.toPlainText())

    docx_page.browser.anchorClicked.emit(QUrl("field:loan_amount"))
    assert dialog.var_list[dialog.table.currentRow()].name == "loan_amount"

    _select(dialog, "name")  # shows the PDF tab
    view = next(v for v in pdf_page.views if "name" in v.rects)
    _select(dialog, "loan_amount")
    center = view.field_rects("name")[0].center().toPoint()
    qtbot.mouseClick(view, Qt.MouseButton.LeftButton, pos=QPoint(center.x(), center.y()))
    assert dialog.var_list[dialog.table.currentRow()].name == "name"


def test_typst_preview_finds_filled_values(qtbot, app):
    source = "#let data = json(bytes(sys.inputs.data))\nГрад #data.city, адрес #data.address, град #data.city\n"
    app.templates.save(
        "addr", "Адрес", TemplateKind.TYPST, {}, {"city": "София", "address": "ул. Витоша 1"}, source_text=source
    )
    dialog = FillDialog(app, ["addr"], {}, "typst-live")
    qtbot.addWidget(dialog)
    dialog.show()
    [page] = dialog.preview.pages
    [view] = page.views
    assert len(view.rects["city"]) == 2 and len(view.rects["address"]) == 1
    xs = sorted(rect[0] for rect in view.rects["city"])
    assert xs[0] < view.rects["address"][0][0] < xs[1]
