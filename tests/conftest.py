import io
import os

import keyring
import pytest
from reportlab.pdfgen import canvas

from blankey.config import load_config
from blankey.core import Blankey

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

PASSWORD = "correct horse battery staple"


@pytest.fixture(autouse=True)
def memory_keyring(monkeypatch):
    store: dict[tuple[str, str], str] = {}
    monkeypatch.setattr(keyring, "get_password", lambda s, u: store.get((s, u)))
    monkeypatch.setattr(keyring, "set_password", lambda s, u, p: store.__setitem__((s, u), p))
    monkeypatch.setattr(keyring, "delete_password", lambda s, u: store.pop((s, u), None))
    return store


@pytest.fixture(autouse=True)
def opened_urls(monkeypatch):
    """Record files the UI would open with the default app instead of launching it."""
    from PySide6.QtGui import QDesktopServices

    opened = []
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()))
    return opened


@pytest.fixture
def app(tmp_path):
    blankey = Blankey(load_config(tmp_path / "data"))
    blankey.vault.initialize(PASSWORD)
    yield blankey
    blankey.vault.close()


@pytest.fixture
def form_pdf() -> bytes:
    """A small AcroForm: name, a 10-cell comb EGN, an email, a checkbox and a radio group."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    form = c.acroForm
    c.drawString(40, 790, "Name")
    form.textfield(name="name", x=100, y=780, width=300, height=18)
    c.drawString(40, 750, "EGN")
    form.textfield(name="egn", x=100, y=740, width=120, height=16, maxlen=10, fieldFlags="comb")
    c.drawString(40, 710, "Email")
    form.textfield(name="email", x=100, y=700, width=80, height=16)
    c.drawString(40, 670, "Bulgarian")
    form.checkbox(name="citizen_bg", x=100, y=665, size=12)
    c.drawString(40, 630, "Status")
    form.radio(name="status", value="single", x=100, y=625, size=12, selected=False)
    form.radio(name="status", value="married", x=140, y=625, size=12, selected=False)
    c.save()
    return buf.getvalue()


@pytest.fixture
def filled_docx(tmp_path):
    """A small filled contract: a bold label run next to plain text, dot blanks and a table."""
    import docx

    document = docx.Document()
    document.add_paragraph("Contract No ......... / ...........2026")
    paragraph = document.add_paragraph()
    paragraph.add_run("1. “ACME HOLDINGS” LTD").bold = True
    paragraph.add_run(", Company ID 999000111, represented by John Smith, ID No .........., card No ..........")
    document.add_paragraph("Loan in the amount of 34 000 EUR.")
    document.add_table(rows=1, cols=2).cell(0, 1).text = "John Smith"
    path = tmp_path / "filled.docx"
    document.save(str(path))
    return path


@pytest.fixture
def flat_pdf() -> bytes:
    """A flat page with printed labels and no form fields."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.drawString(40, 790, "Name:")
    c.drawString(40, 750, "EGN:")
    c.drawString(40, 710, "Married")
    c.save()
    return buf.getvalue()
