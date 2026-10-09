"""Seed a throwaway vault with invented data for screenshots and the demo video.

    uv run python scripts/demo_seed.py /tmp/blankey-demo            # vault + raw contract for Claude to template
    uv run python scripts/demo_seed.py /tmp/blankey-demo --full     # also the contract template + a pending fill

Then run the app against it: BLANKEY_DATA_DIR=/tmp/blankey-demo/data uv run blankey (password: demo).
The raw documents land in <dir>/docs. The target directory must not exist yet.
"""

import argparse
import sys
from pathlib import Path

import docx

from blankey.config import load_config
from blankey.core import Blankey, RequestKind
from blankey.templates import TemplateKind
from blankey.templates import docx as docx_template
from blankey.vault import FieldInput, FieldType

PASSWORD = "demo"

COMPANY = [
    FieldInput("name", "Northwind Studio Ltd", "Company name"),
    FieldInput("company_no", "14839201", "Company number"),
    FieldInput("address", "12 Harbour Street, Bristol BS1 4RN", "Registered office"),
    FieldInput("representative", "Mark Hollis", "Represented by"),
    FieldInput("email", "accounts@northwind.example", "Email", FieldType.EMAIL),
]
PERSON = [
    FieldInput("name", "Jane Example", "Full name"),
    FieldInput("id_number", "X4829173", "ID number"),
    FieldInput("address", "48 Elm Road, Cardiff CF10 3AT", "Home address"),
    FieldInput("iban", "GB33BUKB20201555555555", "IBAN", FieldType.IBAN),
    FieldInput("email", "jane@example.com", "Email", FieldType.EMAIL),
    FieldInput("phone", "+44 7700 900123", "Mobile", FieldType.PHONE),
    FieldInput("birth_date", "14.03.1988", "Date of birth", FieldType.DATE),
]

AGREEMENT_REPLACEMENTS = [
    {"find": "Northwind Studio Ltd", "var": "client_name"},
    {"find": "14839201", "var": "client_company_no"},
    {"find": "12 Harbour Street, Bristol BS1 4RN", "var": "client_address"},
    {"find": "Mark Hollis", "var": "client_representative"},
    {"find": "Jane Example", "var": "contractor_name"},
    {"find": "..............................", "var": "contractor_iban"},
    {"find": "...................", "var": "contractor_address"},
    {"find": "............", "var": "agreement_date"},
    {"find": "..........", "var": "contractor_id_number"},
    {"find": "12 000", "var": "fee"},
    {"find": "......", "var": "agreement_no"},
]
AGREEMENT_FIELDS = {
    "client_name": "{{ client.name }}",
    "client_company_no": "{{ client.company_no }}",
    "client_address": "{{ client.address }}",
    "client_representative": "{{ client.representative }}",
    "contractor_name": "{{ contractor.name }}",
    "contractor_id_number": "{{ contractor.id_number }}",
    "contractor_address": "{{ contractor.address }}",
    "contractor_iban": "{{ contractor.iban }}",
    "agreement_date": "{{ today() }}",
}
AGREEMENT_LABELS = {
    "agreement_no": "Agreement number",
    "agreement_date": "Agreement date",
    "client_name": "Client",
    "client_company_no": "Client company number",
    "client_address": "Client registered office",
    "client_representative": "Client representative",
    "contractor_name": "Contractor",
    "contractor_id_number": "Contractor ID number",
    "contractor_address": "Contractor address",
    "fee": "Fee (EUR)",
    "contractor_iban": "Contractor IBAN",
}
ROLES = {"client": "The company ordering the work", "contractor": "The freelancer doing the work"}


def write_agreement(path: Path) -> None:
    """A previously used contract: filled names, dot blanks for the personal details, a signature table."""
    document = docx.Document()
    document.add_heading("SERVICE AGREEMENT", level=1)
    document.add_paragraph("No. ...... dated ............")
    client = document.add_paragraph()
    client.add_run("Northwind Studio Ltd").bold = True
    client.add_run(", company no. 14839201, registered office 12 Harbour Street, Bristol BS1 4RN,")
    client.add_run(" represented by Mark Hollis (the Client), and")
    contractor = document.add_paragraph()
    contractor.add_run("Jane Example").bold = True
    contractor.add_run(", ID no. .........., residing at ..................., (the Contractor).")
    document.add_paragraph("1. Scope. The Contractor designs and delivers a brand identity package for the Client.")
    document.add_paragraph(
        "2. Fee. The Client pays the Contractor EUR 12 000 within 30 days of invoice, "
        "to IBAN .............................."
    )
    document.add_paragraph("3. Confidentiality. Each party keeps the other party's information confidential.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "For the Client", "The Contractor"
    table.cell(1, 0).text, table.cell(1, 1).text = "Mark Hollis", "Jane Example"
    document.save(str(path))


def write_nda(path: Path) -> None:
    document = docx.Document()
    document.add_heading("MUTUAL NON-DISCLOSURE AGREEMENT", level=1)
    parties = document.add_paragraph()
    parties.add_run("Northwind Studio Ltd").bold = True
    parties.add_run(", company no. 14839201, and Jane Example, ID no. X4829173, agree as follows.")
    document.add_paragraph("1. Confidential information is anything marked or reasonably understood as confidential.")
    document.add_paragraph("2. This agreement lasts three years from 01.09.2026.")
    document.save(str(path))


def save_docx(app: Blankey, template_id: str, name: str, source: Path, replacements, fields, labels=None) -> None:
    target = source.with_name(f"{template_id}.template.docx")
    docx_template.tokenize(source, target, replacements)
    app.templates.save(template_id, name, TemplateKind.DOCX, ROLES, fields, source_file=target, labels=labels)
    target.unlink()


def seed(root: Path, full: bool) -> None:
    docs = root / "docs"
    docs.mkdir(parents=True)
    agreement, nda = docs / "service-agreement.docx", docs / "nda.docx"
    write_agreement(agreement)
    write_nda(nda)

    app = Blankey(load_config(root / "data"))
    app.vault.initialize(PASSWORD)
    company = app.vault.create_profile("Northwind Studio", "company")
    app.vault.set_values(company, COMPANY)
    person = app.vault.create_profile("Jane Example", "person")
    app.vault.set_values(person, PERSON)
    profiles = {"client": company, "contractor": person}

    save_docx(
        app,
        "nda",
        "Mutual NDA",
        nda,
        [
            {"find": "Northwind Studio Ltd", "var": "client_name"},
            {"find": "14839201", "var": "client_company_no"},
            {"find": "Jane Example", "var": "contractor_name"},
            {"find": "X4829173", "var": "contractor_id_number"},
            {"find": "01.09.2026", "var": "start_date"},
        ],
        {
            k: AGREEMENT_FIELDS[k]
            for k in ("client_name", "client_company_no", "contractor_name", "contractor_id_number")
        }
        | {"start_date": "{{ today() }}"},
    )
    bindings = {
        "client_name": {"path": "client.name"},
        "client_company_no": {"path": "client.company_no"},
        "contractor_name": {"path": "contractor.name"},
        "contractor_id_number": {"path": "contractor.id_number"},
        "start_date": {"default": True},
    }
    fill_set = app.vault.save_fill_set("Northwind NDA", ["nda"], profiles, bindings)
    document_ids = app.generate_fill(fill_set)

    vault = app.vault
    asked = vault.create_request(
        RequestKind.PROFILE_INPUT, {"fields": [{"key": "iban"}, {"key": "phone"}], "reason": "Payment details"}
    )
    vault.resolve_request(
        asked, "done", {"fields": [{"key": "iban", "status": "saved"}, {"key": "phone", "status": "saved"}]}
    )
    filled = vault.create_request(
        RequestKind.FILL, {"templates": ["nda"], "name": "Northwind NDA", "reason": "Fill the NDA"}
    )
    vault.resolve_request(
        filled, "done", {"outcome": "generated", "fill_set_id": fill_set, "document_ids": document_ids, "empty": []}
    )

    if full:
        save_docx(
            app,
            "service-agreement",
            "Service agreement",
            agreement,
            AGREEMENT_REPLACEMENTS,
            AGREEMENT_FIELDS,
            AGREEMENT_LABELS,
        )
        app.create_request(
            RequestKind.FILL,
            {
                "templates": ["service-agreement"],
                "profiles": profiles,
                "name": "Northwind brand identity",
                "values": {"fee": "12 000", "agreement_no": "2026-014"},
                "paths": {},
                "fill_set_id": None,
                "reason": "New project for Northwind Studio, prepared from your last contract",
            },
        )
    vault.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("root", type=Path)
    parser.add_argument("--full", action="store_true", help="also save the contract template and a pending fill")
    args = parser.parse_args()
    root = args.root.expanduser().resolve()
    if root.exists():
        sys.exit(f"{root} exists; pick a new directory")
    seed(root, args.full)
    print(f"Seeded {root}\n  BLANKEY_DATA_DIR={root / 'data'} uv run blankey   (password: {PASSWORD})")


if __name__ == "__main__":
    main()
